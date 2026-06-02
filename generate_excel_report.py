"""Generate a comprehensive Excel workbook tracking all rules in the database."""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent
DB_META = PROJECT_ROOT / "database" / "metadata"
OUTPUT = PROJECT_ROOT / "database" / "EVTX_Wazuh_Rules_Tracker.xlsx"

TACTIC_ORDER = [
    "initial_access",
    "execution",
    "persistence",
    "privilege_escalation",
    "defense_evasion",
    "credential_access",
    "discovery",
    "lateral_movement",
    "collection",
    "command_and_control",
    "exfiltration",
    "impact",
]

# Colors
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
INCONC_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
STAT_HEADER_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
STAT_HEADER_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="1F4E79")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")

TACTIC_COLORS = {
    "initial_access": "E74C3C",
    "execution": "E67E22",
    "persistence": "F39C12",
    "privilege_escalation": "27AE60",
    "defense_evasion": "2ECC71",
    "credential_access": "1ABC9C",
    "discovery": "3498DB",
    "lateral_movement": "2980B9",
    "collection": "9B59B6",
    "command_and_control": "8E44AD",
    "exfiltration": "E91E63",
    "impact": "C0392B",
}

SOURCE_COLORS = {
    "sysmon": "3498DB",
    "security": "E74C3C",
    "powershell": "2ECC71",
    "system": "F39C12",
    "other": "95A5A6",
}


def _style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def _autofit(ws, min_width=10, max_width=50):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        lengths = []
        for cell in col_cells:
            if cell.value:
                val = str(cell.value)
                line_max = max((len(line) for line in val.split("\n")), default=0)
                lengths.append(line_max)
        width = max(lengths) + 2 if lengths else min_width
        ws.column_dimensions[col_letter].width = max(min_width, min(width, max_width))


def _load_data():
    index = json.load(open(DB_META / "rule_index.json"))
    validation = json.load(open(DB_META / "validation_results.json"))
    allocations = json.load(open(DB_META / "id_allocations.json"))
    errors = json.load(open(DB_META / "sigma_conversion_errors.json"))
    return index, validation, allocations, errors


def _build_all_rules(wb, index, validation):
    ws = wb.active
    ws.title = "All Rules"
    ws.sheet_properties.tabColor = "1F4E79"

    headers = [
        "Rule ID",
        "Level",
        "Tactic",
        "Technique ID",
        "Technique Name",
        "MITRE IDs",
        "Source Category",
        "Parent SID",
        "Confidence",
        "Origin",
        "Sigma ID",
        "Sigma Level",
        "Tested",
        "Test Result",
        "Test Mode",
        "Deployed",
        "Field Matches",
        "Source File",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    _style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    row = 2
    for rid in sorted(index, key=lambda x: int(x)):
        meta = index[rid]
        val = validation.get(rid, {})
        is_sigma = bool(meta.get("sigma_id"))
        origin = "Sigma" if is_sigma else "EVTX"

        tested = "Yes" if rid in validation else "No"
        if val.get("inconclusive"):
            test_result = "Inconclusive"
        elif val.get("passed"):
            test_result = "PASS"
        elif rid in validation:
            test_result = "FAIL"
        else:
            test_result = ""

        mitre_ids = ", ".join(meta.get("mitre_ids", []))
        field_str = "; ".join(f"{k}={v}" for k, v in meta.get("field_matches", {}).items())
        source_file = meta.get("source_evtx", "")
        if source_file:
            source_file = Path(source_file).name

        values = [
            int(rid),
            meta.get("level", 0),
            meta.get("tactic", ""),
            meta.get("technique_id", ""),
            meta.get("technique_name", ""),
            mitre_ids,
            meta.get("source_category", ""),
            meta.get("parent_sid", ""),
            meta.get("confidence", ""),
            origin,
            meta.get("sigma_id", ""),
            meta.get("sigma_level", ""),
            tested,
            test_result,
            val.get("mode", ""),
            "No",
            field_str,
            source_file,
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top")

        # Color test result
        result_cell = ws.cell(row=row, column=14)
        if test_result == "PASS":
            result_cell.fill = PASS_FILL
        elif test_result == "FAIL":
            result_cell.fill = FAIL_FILL
        elif test_result == "Inconclusive":
            result_cell.fill = INCONC_FILL

        row += 1

    _autofit(ws)
    ws.column_dimensions["Q"].width = 60  # field matches
    ws.column_dimensions["R"].width = 40  # source file
    return row - 2


def _build_stats(wb, index, validation, allocations, errors):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_properties.tabColor = "2E75B6"

    # ── Overview ──
    ws.merge_cells("A1:F1")
    cell = ws.cell(row=1, column=1, value="EVTX-Wazuh Rule Database Dashboard")
    cell.font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    cell.alignment = CENTER

    total = len(index)
    sigma_count = sum(1 for m in index.values() if m.get("sigma_id"))
    evtx_count = total - sigma_count
    passed = sum(1 for v in validation.values() if v.get("passed"))
    failed = sum(1 for v in validation.values() if not v.get("passed") and not v.get("inconclusive"))
    inconclusive = sum(1 for v in validation.values() if v.get("inconclusive"))

    # Summary box
    summary = [
        ("Total Rules", total),
        ("EVTX-Generated", evtx_count),
        ("Sigma-Converted", sigma_count),
        ("", ""),
        ("Tested", len(validation)),
        ("Passed", passed),
        ("Failed", failed),
        ("Inconclusive", inconclusive),
        ("Pass Rate (overall)", f"{100 * passed / total:.1f}%"),
        ("Pass Rate (EVTX)", ""),
        ("Pass Rate (Sigma)", ""),
        ("", ""),
        ("Validation Errors", 0),
        ("Sigma Conversion Errors", sum(errors.get("summary", {}).values())),
        ("Deployed", 0),
    ]

    # Compute EVTX/Sigma pass rates
    evtx_pass = sum(1 for rid, v in validation.items() if v.get("passed") and not index.get(rid, {}).get("sigma_id"))
    sigma_pass = sum(1 for rid, v in validation.items() if v.get("passed") and index.get(rid, {}).get("sigma_id"))
    evtx_tested = sum(1 for rid in validation if not index.get(rid, {}).get("sigma_id"))
    sigma_tested = sum(1 for rid in validation if index.get(rid, {}).get("sigma_id"))
    summary[9] = ("Pass Rate (EVTX)", f"{100 * evtx_pass / evtx_tested:.1f}%" if evtx_tested else "N/A")
    summary[10] = ("Pass Rate (Sigma)", f"{100 * sigma_pass / sigma_tested:.1f}%" if sigma_tested else "N/A")

    row = 3
    ws.cell(row=row, column=1, value="Overview").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.cell(row=row, column=2).fill = SECTION_FILL
    row += 1
    for label, val in summary:
        if not label:
            row += 1
            continue
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=val)
        ws.cell(row=row, column=1).border = THIN_BORDER
        ws.cell(row=row, column=2).border = THIN_BORDER
        row += 1

    # ── Tactic Distribution ──
    row += 2
    tactic_start = row
    ws.cell(row=row, column=1, value="Rules by Tactic").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    for c in range(2, 5):
        ws.cell(row=row, column=c).fill = SECTION_FILL
    row += 1
    ws.cell(row=row, column=1, value="Tactic")
    ws.cell(row=row, column=2, value="Count")
    ws.cell(row=row, column=3, value="Capacity")
    ws.cell(row=row, column=4, value="% Used")
    _style_header(ws, row, 4)
    row += 1
    tactic_data_start = row

    by_tactic = {}
    for meta in index.values():
        t = meta.get("tactic", "other")
        by_tactic[t] = by_tactic.get(t, 0) + 1

    for tactic in TACTIC_ORDER:
        count = by_tactic.get(tactic, 0)
        info = allocations.get(tactic, {})
        cap = info.get("range_end", 0) - info.get("range_start", 0) + 1 if info else 1000
        pct = count / cap if cap else 0
        ws.cell(row=row, column=1, value=tactic).border = THIN_BORDER
        ws.cell(row=row, column=2, value=count).border = THIN_BORDER
        ws.cell(row=row, column=3, value=cap).border = THIN_BORDER
        pct_cell = ws.cell(row=row, column=4, value=pct)
        pct_cell.number_format = "0.0%"
        pct_cell.border = THIN_BORDER
        if pct >= 1.0:
            pct_cell.fill = FAIL_FILL
        elif pct >= 0.8:
            pct_cell.fill = INCONC_FILL
        row += 1

    # Tactic bar chart
    chart = BarChart()
    chart.type = "col"
    chart.title = "Rules by MITRE ATT&CK Tactic"
    chart.y_axis.title = "Rule Count"
    chart.x_axis.title = "Tactic"
    chart.style = 10
    chart.width = 24
    chart.height = 14
    data_ref = Reference(
        ws, min_col=2, min_row=tactic_data_start - 1, max_row=tactic_data_start + len(TACTIC_ORDER) - 1
    )
    cats_ref = Reference(ws, min_col=1, min_row=tactic_data_start, max_row=tactic_data_start + len(TACTIC_ORDER) - 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.shape = 4
    ws.add_chart(chart, f"F{tactic_start}")

    # ── Source Distribution ──
    row += 2
    source_start = row
    ws.cell(row=row, column=1, value="Rules by Source").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.cell(row=row, column=2).fill = SECTION_FILL
    row += 1
    ws.cell(row=row, column=1, value="Source")
    ws.cell(row=row, column=2, value="Count")
    _style_header(ws, row, 2)
    row += 1
    source_data_start = row

    by_source = {}
    for meta in index.values():
        s = meta.get("source_category", "other")
        by_source[s] = by_source.get(s, 0) + 1

    for source in ["sysmon", "security", "powershell", "system", "other"]:
        count = by_source.get(source, 0)
        if count == 0:
            continue
        ws.cell(row=row, column=1, value=source).border = THIN_BORDER
        ws.cell(row=row, column=2, value=count).border = THIN_BORDER
        row += 1

    # Source pie chart
    pie = PieChart()
    pie.title = "Rules by Event Source"
    pie.style = 10
    pie.width = 16
    pie.height = 12
    pie_data = Reference(ws, min_col=2, min_row=source_data_start - 1, max_row=row - 1)
    pie_cats = Reference(ws, min_col=1, min_row=source_data_start, max_row=row - 1)
    pie.add_data(pie_data, titles_from_data=True)
    pie.set_categories(pie_cats)
    ws.add_chart(pie, f"F{source_start}")

    # ── Confidence Distribution ──
    row += 2
    ws.cell(row=row, column=1, value="Rules by Confidence").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.cell(row=row, column=2).fill = SECTION_FILL
    row += 1
    ws.cell(row=row, column=1, value="Confidence")
    ws.cell(row=row, column=2, value="Count")
    _style_header(ws, row, 2)
    row += 1

    by_conf = {}
    for meta in index.values():
        c = meta.get("confidence", "unknown")
        by_conf[c] = by_conf.get(c, 0) + 1

    for conf in ["high", "medium", "low", "unknown"]:
        count = by_conf.get(conf, 0)
        if count:
            ws.cell(row=row, column=1, value=conf).border = THIN_BORDER
            ws.cell(row=row, column=2, value=count).border = THIN_BORDER
            row += 1

    # ── Level Distribution ──
    row += 2
    ws.cell(row=row, column=1, value="Rules by Severity Level").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.cell(row=row, column=2).fill = SECTION_FILL
    row += 1
    ws.cell(row=row, column=1, value="Level")
    ws.cell(row=row, column=2, value="Count")
    _style_header(ws, row, 2)
    row += 1

    by_level = {}
    for meta in index.values():
        lv = meta.get("level", 0)
        by_level[lv] = by_level.get(lv, 0) + 1

    for lv in sorted(by_level):
        ws.cell(row=row, column=1, value=lv).border = THIN_BORDER
        ws.cell(row=row, column=2, value=by_level[lv]).border = THIN_BORDER
        row += 1

    # ── Sigma Error Summary ──
    row += 2
    ws.cell(row=row, column=1, value="Sigma Conversion Errors").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.cell(row=row, column=2).fill = SECTION_FILL
    row += 1
    ws.cell(row=row, column=1, value="Category")
    ws.cell(row=row, column=2, value="Count")
    _style_header(ws, row, 2)
    row += 1

    for cat, cnt in sorted(errors.get("summary", {}).items(), key=lambda x: -x[1]):
        ws.cell(row=row, column=1, value=cat).border = THIN_BORDER
        ws.cell(row=row, column=2, value=cnt).border = THIN_BORDER
        row += 1

    _autofit(ws, min_width=14)
    ws.column_dimensions["A"].width = 30


def _build_validation(wb, index, validation):
    ws = wb.create_sheet("Validation Results")
    ws.sheet_properties.tabColor = "27AE60"

    headers = [
        "Rule ID",
        "Tactic",
        "Technique",
        "Test Result",
        "Test Mode",
        "Confidence",
        "Origin",
        "Error",
        "Details",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    _style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    row = 2
    for rid in sorted(validation, key=lambda x: int(x)):
        val = validation[rid]
        meta = index.get(rid, {})
        is_sigma = bool(meta.get("sigma_id"))

        if val.get("inconclusive"):
            result = "Inconclusive"
        elif val.get("passed"):
            result = "PASS"
        else:
            result = "FAIL"

        details = "\n".join(val.get("details", []))

        values = [
            int(rid),
            meta.get("tactic", ""),
            meta.get("technique_id", ""),
            result,
            val.get("mode", ""),
            meta.get("confidence", ""),
            "Sigma" if is_sigma else "EVTX",
            val.get("error", ""),
            details,
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top")

        result_cell = ws.cell(row=row, column=4)
        if result == "PASS":
            result_cell.fill = PASS_FILL
        elif result == "FAIL":
            result_cell.fill = FAIL_FILL
        else:
            result_cell.fill = INCONC_FILL

        row += 1

    _autofit(ws)
    ws.column_dimensions["I"].width = 80


def _build_by_tactic(wb, index, validation):
    for tactic in TACTIC_ORDER:
        rules = {rid: meta for rid, meta in index.items() if meta.get("tactic") == tactic}
        if not rules:
            continue

        name = tactic.replace("_", " ").title()
        ws = wb.create_sheet(name[:31])
        color = TACTIC_COLORS.get(tactic, "808080")
        ws.sheet_properties.tabColor = color

        headers = [
            "Rule ID",
            "Level",
            "Technique ID",
            "Technique Name",
            "Confidence",
            "Origin",
            "Test Result",
            "MITRE IDs",
            "Field Matches",
            "Source File",
        ]
        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
        _style_header(ws, 1, len(headers))
        ws.freeze_panes = "A2"

        row = 2
        for rid in sorted(rules, key=lambda x: int(x)):
            meta = rules[rid]
            val = validation.get(rid, {})
            is_sigma = bool(meta.get("sigma_id"))

            if val.get("inconclusive"):
                result = "Inconclusive"
            elif val.get("passed"):
                result = "PASS"
            elif rid in validation:
                result = "FAIL"
            else:
                result = "Not tested"

            mitre_ids = ", ".join(meta.get("mitre_ids", []))
            field_str = "; ".join(f"{k}={v}" for k, v in meta.get("field_matches", {}).items())
            source_file = Path(meta.get("source_evtx", "")).name if meta.get("source_evtx") else ""

            values = [
                int(rid),
                meta.get("level", 0),
                meta.get("technique_id", ""),
                meta.get("technique_name", ""),
                meta.get("confidence", ""),
                "Sigma" if is_sigma else "EVTX",
                result,
                mitre_ids,
                field_str,
                source_file,
            ]
            for col, v in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=v)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="top")

            result_cell = ws.cell(row=row, column=7)
            if result == "PASS":
                result_cell.fill = PASS_FILL
            elif result == "FAIL":
                result_cell.fill = FAIL_FILL
            elif result == "Inconclusive":
                result_cell.fill = INCONC_FILL

            row += 1

        _autofit(ws)
        ws.column_dimensions["I"].width = 60


def _build_sigma_errors(wb, errors):
    ws = wb.create_sheet("Sigma Errors")
    ws.sheet_properties.tabColor = "E74C3C"

    headers = ["File", "Sigma ID", "Error Category", "Message"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    _style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    row = 2
    for rec in errors.get("records", []):
        fname = Path(rec.get("file", "")).name
        values = [
            fname,
            rec.get("sigma_id", ""),
            rec.get("category", ""),
            rec.get("message", ""),
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = THIN_BORDER
        row += 1

    _autofit(ws)
    ws.column_dimensions["D"].width = 80


def _build_deployment(wb, index, validation):
    ws = wb.create_sheet("Deployment Tracker")
    ws.sheet_properties.tabColor = "F39C12"

    headers = [
        "Rule ID",
        "Tactic",
        "Technique ID",
        "Level",
        "Confidence",
        "Origin",
        "Test Result",
        "Deployed",
        "Deploy Date",
        "Deploy Target",
        "Notes",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    _style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    row = 2
    for rid in sorted(index, key=lambda x: int(x)):
        meta = index[rid]
        val = validation.get(rid, {})
        is_sigma = bool(meta.get("sigma_id"))

        if val.get("passed"):
            result = "PASS"
        elif rid in validation and not val.get("inconclusive"):
            result = "FAIL"
        elif val.get("inconclusive"):
            result = "Inconclusive"
        else:
            result = ""

        values = [
            int(rid),
            meta.get("tactic", ""),
            meta.get("technique_id", ""),
            meta.get("level", 0),
            meta.get("confidence", ""),
            "Sigma" if is_sigma else "EVTX",
            result,
            "No",
            "",
            "",
            "",
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = THIN_BORDER

        row += 1

    _autofit(ws)
    ws.column_dimensions["K"].width = 40


def main():
    print("Loading metadata...")
    index, validation, allocations, errors = _load_data()

    wb = Workbook()

    print("Building All Rules sheet...")
    count = _build_all_rules(wb, index, validation)

    print("Building Dashboard...")
    _build_stats(wb, index, validation, allocations, errors)

    print("Building Validation Results...")
    _build_validation(wb, index, validation)

    print("Building per-tactic sheets...")
    _build_by_tactic(wb, index, validation)

    print("Building Sigma Errors...")
    _build_sigma_errors(wb, errors)

    print("Building Deployment Tracker...")
    _build_deployment(wb, index, validation)

    # Move Dashboard to first position
    dashboard_idx = wb.sheetnames.index("Dashboard")
    wb.move_sheet("Dashboard", offset=-dashboard_idx)

    print(f"Saving to {OUTPUT}...")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"Done! {count} rules across {len(wb.sheetnames)} sheets.")
    print(f"Sheets: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
